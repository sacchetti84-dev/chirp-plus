# Copyright 2026 CHIRP Plus contributors
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <http://www.gnu.org/licenses/>.

"""Excel spreadsheets with the same columns as the generic CSV format.

.xlsx files are read and written (openpyxl); old .xls files (Excel 97-2003)
can only be read (xlrd). Both libraries are imported only when a file is
actually used, so CHIRP keeps working without them.
"""

import logging

from chirp import directory, errors
from chirp.drivers import generic_csv

LOG = logging.getLogger(__name__)

SHEET_NAME = 'CHIRP'
XLSX_MAGIC = b'PK\x03\x04'
XLS_MAGIC = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'

# Columns written as numbers instead of text, with their Excel format
NUMERIC_COLUMNS = {
    'Location': (int, '0'),
    'Frequency': (float, '0.000000'),
    'Offset': (float, '0.000000'),
    'rToneFreq': (float, '0.0'),
    'cToneFreq': (float, '0.0'),
    'DtcsCode': (int, '000'),
    'RxDtcsCode': (int, '000'),
    'TStep': (float, '0.00'),
}
COLUMN_WIDTHS = {'Name': 14, 'Frequency': 13, 'Offset': 11, 'Comment': 30,
                 'CrossMode': 12, 'DtcsPolarity': 12, 'Power': 9}


def cell_to_text(value):
    """Turn a spreadsheet cell into the string the CSV parser expects"""
    if value is None:
        return ''
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            # Spreadsheets store 1 as 1.0; "1.0" is not a valid Location
            return str(int(value))
        return repr(value)
    return str(value).strip()


def rows_to_text(rows):
    """Normalize spreadsheet rows for generic_csv's row parser.

    Drops blank rows, trailing empty header cells (and the matching
    columns), so a stray note far to the right does not break the import.
    """
    width = None
    for row in rows:
        row = [cell_to_text(v) for v in row]
        if not any(row):
            continue
        if row[0].startswith('#'):
            # Comment lines: only the text, not the empty cells after it
            yield [cell for cell in row if cell]
            continue
        if width is None:
            width = len(row)
            while width and not row[width - 1]:
                width -= 1
        yield row[:width]


def _number(column, text):
    conv, fmt = NUMERIC_COLUMNS[column]
    try:
        return conv(text), fmt
    except ValueError:
        return text, None


@directory.register
class ExcelRadio(generic_csv.CSVRadio):
    """Excel workbook (.xlsx) in the generic CSV layout"""
    VENDOR = 'Generic'
    MODEL = 'Excel'
    FILE_EXTENSION = 'xlsx'
    FORMATS = [directory.register_format('Excel', '*.xlsx')]

    def load(self, filename=None):
        if filename is None and self._filename is None:
            raise errors.RadioError('Need a location to load from')
        if filename:
            self._filename = filename
        self._blank()
        return self._load_rows(rows_to_text(self._read_rows()))

    def load_from(self, string):
        raise errors.RadioError('Excel files can not be loaded from text')

    def _read_rows(self):
        try:
            import openpyxl
        except ImportError:
            raise errors.RadioError('Reading Excel files needs openpyxl')
        wb = openpyxl.load_workbook(self._filename, read_only=True,
                                    data_only=True)
        try:
            sheet = (wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames
                     else wb.worksheets[0])
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            wb.close()

    def save(self, filename=None):
        if filename is None and self._filename is None:
            raise errors.RadioError('Need a location to save to')
        if filename:
            self._filename = filename
        if not self._filename.lower().endswith('.xlsx'):
            raise errors.RadioError(
                'Excel files can only be saved as .xlsx')
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise errors.RadioError('Writing Excel files needs openpyxl')

        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = SHEET_NAME
        header = None
        for row in self._rows():
            if header is None and row and row[0] == 'Location':
                header = row
                sheet.append(row)
                continue
            if header is None or row[0].startswith('#'):
                # Comment lines are kept as a single text cell
                sheet.append(row)
                continue
            sheet.append([None] * len(row))
            for col, (name, text) in enumerate(zip(header, row), start=1):
                cell = sheet.cell(row=sheet.max_row, column=col)
                if name in NUMERIC_COLUMNS and text != '':
                    cell.value, fmt = _number(name, text)
                    if fmt:
                        cell.number_format = fmt
                else:
                    cell.value = text if text != '' else None

        if header:
            header_row = next(r for r in range(1, sheet.max_row + 1)
                              if sheet.cell(row=r, column=1).value ==
                              'Location')
            for col, name in enumerate(header, start=1):
                cell = sheet.cell(row=header_row, column=col)
                cell.font = Font(bold=True, color='C8FF3D')
                cell.fill = PatternFill('solid', fgColor='1F4D1A')
                cell.alignment = Alignment(horizontal='center')
                sheet.column_dimensions[get_column_letter(col)].width = (
                    COLUMN_WIDTHS.get(name, max(8, len(name) + 2)))
            sheet.freeze_panes = sheet.cell(row=header_row + 1, column=2)
            sheet.auto_filter.ref = '%s:%s%i' % (
                sheet.cell(row=header_row, column=1).coordinate,
                get_column_letter(len(header)), sheet.max_row)
        wb.save(self._filename)

    def as_string(self):
        raise errors.RadioError('Excel files have no text representation')

    @classmethod
    def match_model(cls, filedata, filename):
        return (filename.lower().endswith('.xlsx') and
                (filedata.startswith(XLSX_MAGIC) or filedata == b''))


@directory.register
class ExcelLegacyRadio(ExcelRadio):
    """Old Excel 97-2003 workbook (.xls): read only, save as .xlsx"""
    MODEL = 'XLS'
    FORMATS = [directory.register_format('Excel 97-2003', '*.xls',
                                         readonly=True),
               'Excel']

    def _read_rows(self):
        try:
            import xlrd
        except ImportError:
            raise errors.RadioError('Reading .xls files needs xlrd')
        book = xlrd.open_workbook(self._filename)
        try:
            sheet = (book.sheet_by_name(SHEET_NAME)
                     if SHEET_NAME in book.sheet_names()
                     else book.sheet_by_index(0))
            return [sheet.row_values(i) for i in range(sheet.nrows)]
        finally:
            book.release_resources()

    def save(self, filename=None):
        target = filename or self._filename or ''
        if target.lower().endswith('.xls'):
            raise errors.RadioError(
                'Excel 97-2003 (.xls) files can only be read: use Save As '
                'and choose the Excel (.xlsx) format')
        return super().save(filename)

    @classmethod
    def match_model(cls, filedata, filename):
        return (filename.lower().endswith('.xls') and
                filedata.startswith(XLS_MAGIC))
